"""
Phase 7: Build Professional PowerPoint Presentation
Music Streaming Dataset Analysis — 8 Core Slides
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import json
import os

# ── Load results ────────────────────────────────────────────────
with open("output/analysis_results.json") as f:
    R = json.load(f)

CHART_DIR = os.path.join("output", "charts")

# ── Color scheme ────────────────────────────────────────────────
DARK_BG     = RGBColor(0x0F, 0x0F, 0x23)    # Very dark navy
CARD_BG     = RGBColor(0x1A, 0x1A, 0x2E)    # Dark card
ACCENT_1    = RGBColor(0x00, 0xB4, 0xD8)    # Teal
ACCENT_2    = RGBColor(0x7B, 0x2C, 0xBF)    # Purple
ACCENT_3    = RGBColor(0xE7, 0x6F, 0x51)    # Coral
ACCENT_4    = RGBColor(0x06, 0xD6, 0xA0)    # Green
ACCENT_5    = RGBColor(0xFF, 0xD1, 0x66)    # Gold
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY  = RGBColor(0xC0, 0xC0, 0xC0)
MED_GRAY    = RGBColor(0x80, 0x80, 0x90)
SOFT_WHITE  = RGBColor(0xE8, 0xE8, 0xF0)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

# ── Helper functions ────────────────────────────────────────────
def set_slide_bg(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, width, height, fill_color, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def add_text_box(slide, left, top, width, height, text, font_size=14, bold=False, color=WHITE, alignment=PP_ALIGN.LEFT, font_name="Segoe UI"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox

def add_multiline_text(slide, left, top, width, height, lines, font_name="Segoe UI"):
    """lines = [(text, size, bold, color, alignment), ...]"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, alignment) in enumerate(lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = font_name
        p.alignment = alignment
        p.space_after = Pt(4)
    return txBox

def add_accent_line(slide, left, top, width, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Pt(3))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

# ================================================================
# SLIDE 1: Title
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
set_slide_bg(slide, DARK_BG)

# Accent line at top
add_accent_line(slide, Inches(0), Inches(0), SLIDE_W, ACCENT_1)

# Title
add_text_box(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.2),
             "Music Streaming Dataset Analysis", 40, True, WHITE, PP_ALIGN.CENTER)

# Subtitle
add_text_box(slide, Inches(1.5), Inches(3.3), Inches(10), Inches(0.8),
             "Understanding the Drivers of Track Popularity", 24, False, ACCENT_1, PP_ALIGN.CENTER)

# Accent line under subtitle
add_accent_line(slide, Inches(5), Inches(4.3), Inches(3.333), ACCENT_1)

# Metadata
add_text_box(slide, Inches(1.5), Inches(5.2), Inches(10), Inches(0.4),
             "Data Analyst Assessment  |  Tech Round 2", 16, False, MED_GRAY, PP_ALIGN.CENTER)

add_text_box(slide, Inches(1.5), Inches(5.7), Inches(10), Inches(0.4),
             f"Dataset: {R['raw_shape'][0]:,} tracks  |  {114} genres  |  5 market regions", 
             14, False, MED_GRAY, PP_ALIGN.CENTER)

# ================================================================
# SLIDE 2: Executive Summary
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(5), Inches(0.6),
             "Executive Summary", 28, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(2.5), ACCENT_1)

# Finding cards
card_data = [
    ("1", "Energy Has No Overall Effect on Popularity", 
     f"r = {R['h1_overall']['pearson_r']:.3f}, R\u00b2 < 0.01% \u2014 energy alone explains virtually none of the variation in popularity. However, the relationship varies substantially by genre (r ranges from \u22120.40 in emo to +0.39 in j-pop).",
     ACCENT_1),
    ("2", "No Regional Bias in Superstar Distribution", 
     f"Superstar artists are evenly spread across all 5 regions (3.9\u20135.3%). Chi-square test: p = 0.785, Cramer's V = 0.017. There is no statistically significant evidence of regional concentration.",
     ACCENT_2),
    ("3", "Artist Tier Is the Strongest Popularity Driver", 
     "Superstars average 76.1 popularity vs 11.9 for Emerging artists \u2014 a 64-point gap. This structural variable dominates all audio features in predicting popularity.",
     ACCENT_4),
]

for i, (num, title, body, accent) in enumerate(card_data):
    y = Inches(1.3) + Inches(i * 1.95)
    card = add_rect(slide, Inches(0.6), y, Inches(12.1), Inches(1.7), CARD_BG, accent)
    
    # Number circle
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.0), y + Inches(0.35), Inches(0.65), Inches(0.65))
    circle.fill.solid()
    circle.fill.fore_color.rgb = accent
    circle.line.fill.background()
    tf = circle.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = num
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    p.alignment = PP_ALIGN.CENTER
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    add_text_box(slide, Inches(2.0), y + Inches(0.2), Inches(10), Inches(0.45),
                 title, 17, True, WHITE)
    add_text_box(slide, Inches(2.0), y + Inches(0.7), Inches(10.2), Inches(0.9),
                 body, 12, False, LIGHT_GRAY)

# Confidence note
add_text_box(slide, Inches(0.6), Inches(7.0), Inches(12), Inches(0.4),
             "Overall Confidence: Moderate to High  |  Findings based on 9,637 tracks after cleaning  |  Association \u2260 Causation",
             11, False, MED_GRAY, PP_ALIGN.LEFT)

# ================================================================
# SLIDE 3: Data Quality & Cleaning
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(5), Inches(0.6),
             "Data Quality & Cleaning", 28, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(2.5), ACCENT_1)

# Left column: Issues Found
add_text_box(slide, Inches(0.6), Inches(1.2), Inches(6), Inches(0.4),
             "Issues Discovered", 18, True, ACCENT_1)

issues = [
    ("\u2022  Casing inconsistencies across 4 columns (market_region, artist_tier, streaming_platform, explicit) \u2014 e.g. 'Europe' vs 'EUROPE', 'Emerging' vs 'emerging'",),
    ("\u2022  250 non-numeric release_year values ('late 90s', 'circa 2010', 'unknown', 'pre-2000') plus 394 missing",),
    ("\u2022  74 exact duplicate rows",),
    ("\u2022  20 invalid popularity scores (negative values, scores > 100)",),
    ("\u2022  4 negative track durations (\u22121000ms)",),
    ("\u2022  114 granular genre labels (same track can appear under multiple genres)",),
]

y_offset = Inches(1.7)
for issue in issues:
    tb = add_text_box(slide, Inches(0.8), y_offset, Inches(5.8), Inches(0.45),
                      issue[0], 11, False, LIGHT_GRAY)
    y_offset += Inches(0.45)

# Right column: Treatment Applied
add_text_box(slide, Inches(6.8), Inches(1.2), Inches(6), Inches(0.4),
             "Treatment Applied", 18, True, ACCENT_4)

treatments = [
    ("\u2022  Standardised all categorical values to title case; fixed 'Amazon  Music' double-space",),
    ("\u2022  Non-numeric years \u2192 NaN; derived clean_decade from valid years only (not from the existing release_decade column)",),
    ("\u2022  Removed 74 exact duplicates; retained non-identical duplicates (same track, different context)",),
    ("\u2022  Invalid popularity \u2192 NaN (outside documented 0\u2013100 range)",),
    ("\u2022  Negative durations \u2192 NaN",),
    ("\u2022  Retained granular genres for analysis; used top genres for visualisation",),
]

y_offset = Inches(1.7)
for treatment in treatments:
    tb = add_text_box(slide, Inches(7.0), y_offset, Inches(5.8), Inches(0.45),
                      treatment[0], 11, False, LIGHT_GRAY)
    y_offset += Inches(0.45)

# Summary stats box
summary_card = add_rect(slide, Inches(0.6), Inches(5.2), Inches(12.1), Inches(1.8), CARD_BG, ACCENT_1)
add_multiline_text(slide, Inches(1.0), Inches(5.3), Inches(11.5), Inches(1.6), [
    ("Cleaning Impact Summary", 14, True, ACCENT_1, PP_ALIGN.LEFT),
    (f"Raw dataset: {R['raw_shape'][0]:,} rows  \u2192  Cleaned: {R['cleaned_shape'][0]:,} rows (74 duplicates removed)", 12, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    (f"Popularity available: 9,637 tracks  |  Valid release years: ~9,400 tracks  |  All 5 regions & 4 tiers fully populated after casing fix", 12, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Design principle: Every exclusion has a documented rationale. Missing values are excluded per analysis, not globally deleted.", 12, False, ACCENT_5, PP_ALIGN.LEFT),
])

# ================================================================
# SLIDE 4: Analytical Approach
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(5), Inches(0.6),
             "Analytical Approach", 28, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(2.5), ACCENT_1)

# Pipeline steps
steps = [
    ("Raw Data", "10,132 rows\n24 columns", ACCENT_3),
    ("Clean", "Fix casing, remove\nduplicates & invalids", ACCENT_1),
    ("Define Outcome", "Popularity (0\u2013100)\nas target variable", ACCENT_2),
    ("Hypotheses", "Energy \u00d7 Popularity\nSuperstar regions", ACCENT_4),
    ("Analyse", "Correlation, chi-square\neffect sizes", ACCENT_5),
    ("Conclude", "Evidence strength\n& limitations", ACCENT_1),
]

step_w = Inches(1.8)
gap = Inches(0.25)
total_w = len(steps) * step_w + (len(steps) - 1) * gap
start_x = (SLIDE_W - total_w) // 2

for i, (title, desc, color) in enumerate(steps):
    x = start_x + i * (step_w + gap)
    y = Inches(1.5)
    
    card = add_rect(slide, x, y, step_w, Inches(1.8), CARD_BG, color)
    
    # Step number
    add_text_box(slide, x, y + Inches(0.1), step_w, Inches(0.35),
                 f"Step {i+1}", 10, False, color, PP_ALIGN.CENTER)
    add_text_box(slide, x, y + Inches(0.4), step_w, Inches(0.4),
                 title, 14, True, WHITE, PP_ALIGN.CENTER)
    add_text_box(slide, x + Inches(0.1), y + Inches(0.9), step_w - Inches(0.2), Inches(0.8),
                 desc, 10, False, LIGHT_GRAY, PP_ALIGN.CENTER)

# Key principles
add_text_box(slide, Inches(0.6), Inches(3.8), Inches(6), Inches(0.4),
             "Key Analytical Principles", 18, True, ACCENT_1)

principles = [
    "\u2022  Popularity is the outcome variable (0\u2013100 scale, higher = more popular at time of data pull)",
    "\u2022  Association \u2260 Causation: observational data cannot establish that audio features cause popularity",
    "\u2022  Statistical significance \u2260 practical importance: with ~10,000 records, tiny effects can be 'significant'",
    "\u2022  Effect sizes and confidence intervals matter more than p-values alone",
    "\u2022  Confounders: genre, artist tier, platform, and playlist context all influence popularity independently",
]

y_start = Inches(4.3)
for j, principle in enumerate(principles):
    add_text_box(slide, Inches(0.8), y_start + j * Inches(0.4), Inches(11.5), Inches(0.4),
                 principle, 12, False, LIGHT_GRAY)

# Hypotheses tested box
hyp_card = add_rect(slide, Inches(0.6), Inches(6.4), Inches(12.1), Inches(0.8), CARD_BG, ACCENT_2)
add_text_box(slide, Inches(1.0), Inches(6.5), Inches(11.5), Inches(0.6),
             "Hypotheses Tested:   H1: Does energy drive popularity? Does this vary by genre?   |   H2: Are Superstar artists concentrated in specific regions?",
             12, False, WHITE, PP_ALIGN.LEFT)

# ================================================================
# SLIDE 5: Hypothesis 1 — Energy vs Popularity
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(8), Inches(0.6),
             "Hypothesis 1: Does Energy Drive Popularity?", 24, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(3), ACCENT_1)

# Chart
chart_path = os.path.join(CHART_DIR, "h1_energy_popularity.png")
slide.shapes.add_picture(chart_path, Inches(0.4), Inches(1.1), Inches(7.2), Inches(4.0))

# Stats card on the right
stats_card = add_rect(slide, Inches(7.8), Inches(1.1), Inches(5.0), Inches(4.0), CARD_BG, ACCENT_1)

h1 = R["h1_overall"]
add_multiline_text(slide, Inches(8.1), Inches(1.2), Inches(4.6), Inches(3.8), [
    ("Overall Relationship", 16, True, ACCENT_1, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    (f"Pearson r = {h1['pearson_r']:.4f}", 14, True, WHITE, PP_ALIGN.LEFT),
    (f"Spearman \u03c1 = {h1['spearman_r']:.4f}", 14, True, WHITE, PP_ALIGN.LEFT),
    (f"R\u00b2 = {h1['r_squared']:.4f}  (<0.01%)", 14, True, WHITE, PP_ALIGN.LEFT),
    (f"p = {h1['pearson_p']:.3f}  (not significant)", 13, False, ACCENT_3, PP_ALIGN.LEFT),
    (f"N = {h1['n']:,}", 13, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Interpretation", 14, True, ACCENT_5, PP_ALIGN.LEFT),
    ("Energy explains essentially zero variance in popularity overall. The relationship is flat, not statistically significant, and has no practical importance.", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# Bottom interpretation
interp_card = add_rect(slide, Inches(0.4), Inches(5.4), Inches(12.5), Inches(1.7), CARD_BG, ACCENT_5)
add_multiline_text(slide, Inches(0.8), Inches(5.5), Inches(11.8), Inches(1.5), [
    ("What did we find?", 14, True, ACCENT_5, PP_ALIGN.LEFT),
    ("Mean popularity is remarkably consistent across all energy levels (~26\u201335). Low-energy and high-energy tracks achieve similar popularity. Energy alone is not a driver of track success. However, the story changes when we look within individual genres...", 12, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# ================================================================
# SLIDE 6: Hypothesis 1 — Genre Differences
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(10), Inches(0.6),
             "Hypothesis 1b: Energy\u2013Popularity Varies Dramatically by Genre", 22, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(3), ACCENT_1)

# Chart
chart_path = os.path.join(CHART_DIR, "h1_genre_correlations.png")
slide.shapes.add_picture(chart_path, Inches(0.3), Inches(1.1), Inches(7.5), Inches(4.5))

# Right panel: key findings
right_card = add_rect(slide, Inches(8.0), Inches(1.1), Inches(4.8), Inches(4.5), CARD_BG, ACCENT_2)

sig_genres = R["h1_significant_genres"]
pos_genres = [g for g in sig_genres if g["pearson_r"] > 0]
neg_genres = [g for g in sig_genres if g["pearson_r"] < 0]

lines = [
    ("24 genres show significant correlations", 14, True, ACCENT_1, PP_ALIGN.LEFT),
    ("(p < 0.05 and |r| > 0.1)", 11, False, MED_GRAY, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Strongest Positive (energy \u2192 \u2191 popularity)", 12, True, ACCENT_4, PP_ALIGN.LEFT),
]
for g in pos_genres[:5]:
    lines.append((f"  {g['genre']}: r = {g['pearson_r']:+.3f} (n={g['n']})", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT))

lines.append(("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT))
lines.append(("Strongest Negative (energy \u2192 \u2193 popularity)", 12, True, ACCENT_3, PP_ALIGN.LEFT))
for g in neg_genres[-5:]:
    lines.append((f"  {g['genre']}: r = {g['pearson_r']:+.3f} (n={g['n']})", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT))

add_multiline_text(slide, Inches(8.3), Inches(1.2), Inches(4.3), Inches(4.3), lines)

# Bottom interpretation
interp_card = add_rect(slide, Inches(0.3), Inches(5.8), Inches(12.7), Inches(1.4), CARD_BG, ACCENT_5)
add_multiline_text(slide, Inches(0.7), Inches(5.9), Inches(12.0), Inches(1.2), [
    ("Key Insight", 14, True, ACCENT_5, PP_ALIGN.LEFT),
    ("The null overall correlation masks a genuine pattern: in genres like j-pop, k-pop, and classical, higher energy is positively associated with popularity. In genres like emo, garage, and indie, the reverse is true. Genre context is essential \u2014 energy's 'effect' depends entirely on the musical context.", 12, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# ================================================================
# SLIDE 7: Hypothesis 2 — Superstar Regional Distribution
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(10), Inches(0.6),
             "Hypothesis 2: Is There Regional Bias in Superstar Distribution?", 22, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(3), ACCENT_1)

# Chart
chart_path = os.path.join(CHART_DIR, "h2_superstar_regions.png")
slide.shapes.add_picture(chart_path, Inches(0.2), Inches(1.1), Inches(8.5), Inches(3.9))

# Stats card on right
stats_card = add_rect(slide, Inches(8.9), Inches(1.1), Inches(4.1), Inches(3.9), CARD_BG, ACCENT_2)

h2 = R["h2_chi2"]
add_multiline_text(slide, Inches(9.2), Inches(1.2), Inches(3.6), Inches(3.7), [
    ("Statistical Test", 15, True, ACCENT_2, PP_ALIGN.LEFT),
    ("Chi-square test of independence", 11, False, MED_GRAY, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    (f"\u03c7\u00b2 = {h2['chi2']:.1f},  df = {h2['dof']}", 13, True, WHITE, PP_ALIGN.LEFT),
    (f"p = {h2['p_value']:.3f}", 13, True, ACCENT_4, PP_ALIGN.LEFT),
    (f"Cramer's V = {h2['cramers_v']:.4f}", 13, True, WHITE, PP_ALIGN.LEFT),
    (f"N = {h2['n']:,}", 13, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Verdict", 14, True, ACCENT_4, PP_ALIGN.LEFT),
    ("NOT significant", 16, True, ACCENT_4, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Superstar artists are distributed proportionally across all regions. No evidence of regional bias.", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# Bottom interpretation
interp_card = add_rect(slide, Inches(0.2), Inches(5.2), Inches(12.9), Inches(2.0), CARD_BG, ACCENT_5)

sp = R["h2_superstar_pct"]
region_lines = " | ".join([f"{r}: {sp[r]['pct']:.1f}%" for r in sorted(sp.keys(), key=lambda x: sp[x]['pct'], reverse=True)])

add_multiline_text(slide, Inches(0.6), Inches(5.3), Inches(12.2), Inches(1.8), [
    ("What did we find?", 14, True, ACCENT_5, PP_ALIGN.LEFT),
    (f"Superstar proportion by region:  {region_lines}", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("The spread (3.9%\u20135.3%) is narrow. With p = 0.785 and Cramer's V near zero, the data provides no evidence that Superstar artists cluster in specific regions. The slight variation is well within random chance. However, note that 'Middle East & Africa' has only 588 tracks \u2014 a smaller sample that increases uncertainty for that region.", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# ================================================================
# SLIDE 8: Conclusions, Confidence & Limitations
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide, DARK_BG)

add_text_box(slide, Inches(0.6), Inches(0.3), Inches(8), Inches(0.6),
             "Conclusions, Confidence & Limitations", 26, True, WHITE)
add_accent_line(slide, Inches(0.6), Inches(0.85), Inches(3), ACCENT_1)

# Left: Conclusions
left_card = add_rect(slide, Inches(0.4), Inches(1.2), Inches(6.2), Inches(3.2), CARD_BG, ACCENT_1)
add_multiline_text(slide, Inches(0.7), Inches(1.3), Inches(5.7), Inches(3.0), [
    ("What Appears Associated with Popularity", 15, True, ACCENT_1, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("\u2022  Artist tier is by far the strongest predictor (Superstar: 76 vs Emerging: 12)", 12, False, WHITE, PP_ALIGN.LEFT),
    ("\u2022  Genre matters: pop-film, k-pop, sad, chill genres rank highest", 12, False, WHITE, PP_ALIGN.LEFT),
    ("\u2022  Energy\u2013popularity relationship depends on genre context", 12, False, WHITE, PP_ALIGN.LEFT),
    ("\u2022  Superstar distribution is even across regions", 12, False, WHITE, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("What We Cannot Conclude", 14, True, ACCENT_3, PP_ALIGN.LEFT),
    ("\u2022  Causation: we cannot say energy causes popularity changes", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("\u2022  Temporal effects: vague years limit decade-level analysis", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("\u2022  Platform effects: data is observational, not experimental", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# Right: Confidence & Limitations
right_card = add_rect(slide, Inches(6.8), Inches(1.2), Inches(6.2), Inches(3.2), CARD_BG, ACCENT_2)
add_multiline_text(slide, Inches(7.1), Inches(1.3), Inches(5.7), Inches(3.0), [
    ("Confidence Assessment", 15, True, ACCENT_2, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("H1 (Energy overall): HIGH confidence in null result", 12, True, ACCENT_4, PP_ALIGN.LEFT),
    ("  r \u2248 0, not significant, N = 9,637", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("H1b (Genre variation): MODERATE confidence", 12, True, ACCENT_5, PP_ALIGN.LEFT),
    ("  24 genres significant, but ~80\u201390 per genre", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("H2 (Regional bias): HIGH confidence in null result", 12, True, ACCENT_4, PP_ALIGN.LEFT),
    ("  p = 0.785, effect size negligible", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Key Limitations", 14, True, ACCENT_3, PP_ALIGN.LEFT),
    ("\u2022  Same track can appear under multiple genres/contexts", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("\u2022  Popularity is a snapshot, not a causal outcome", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("\u2022  Multiple testing: 114 genre correlations increase false positive risk", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# Bottom: Artist tier chart
chart_path = os.path.join(CHART_DIR, "popularity_by_tier.png")
slide.shapes.add_picture(chart_path, Inches(0.4), Inches(4.6), Inches(6.0), Inches(2.7))

# Bottom-right: Final takeaway
final_card = add_rect(slide, Inches(6.8), Inches(4.6), Inches(6.2), Inches(2.7), CARD_BG, ACCENT_5)
add_multiline_text(slide, Inches(7.1), Inches(4.7), Inches(5.7), Inches(2.5), [
    ("The Biggest Story in the Data", 15, True, ACCENT_5, PP_ALIGN.LEFT),
    ("", 4, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("Audio features like energy are poor predictors of popularity in isolation. The strongest driver by far is artist tier \u2014 who you are matters far more than how your music sounds.", 13, False, WHITE, PP_ALIGN.LEFT),
    ("", 6, False, LIGHT_GRAY, PP_ALIGN.LEFT),
    ("This finding is consistent with music industry research: popularity is driven by marketing, branding, and platform placement rather than acoustic properties alone.", 11, False, LIGHT_GRAY, PP_ALIGN.LEFT),
])

# ================================================================
# Save
# ================================================================
output_path = "Music_Streaming_Analysis.pptx"
prs.save(output_path)
print(f"[DONE] PowerPoint saved to: {output_path}")
print(f"  Slides: {len(prs.slides)}")
