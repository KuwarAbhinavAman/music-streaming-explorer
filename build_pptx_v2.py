"""
Revised PowerPoint — Light, Clean, Professional
Admiral-inspired palette | Plain English | Business-first messaging
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import json
import os

# ── Load results ────────────────────────────────────────────────
with open("output/analysis_results.json") as f:
    R = json.load(f)

CHART_DIR = os.path.join("output", "charts_v2")

# ── Colour palette ──────────────────────────────────────────────
PRIMARY_BLUE  = RGBColor(0x17, 0x69, 0xAA)
DEEP_BLUE     = RGBColor(0x17, 0x4A, 0x73)
SOFT_BLUE     = RGBColor(0xEA, 0xF4, 0xFA)
VLIGHT_BLUE   = RGBColor(0xF5, 0xF9, 0xFC)
MAGENTA       = RGBColor(0xB8, 0x3A, 0x7A)
SOFT_MAGENTA  = RGBColor(0xF8, 0xED, 0xF4)
VLIGHT_PINK   = RGBColor(0xFC, 0xF4, 0xF8)
CHARCOAL      = RGBColor(0x27, 0x34, 0x3D)
MUTED_GREY    = RGBColor(0x6B, 0x72, 0x80)
LIGHT_BORDER  = RGBColor(0xDC, 0xE6, 0xEC)
WHITE         = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SLIDE_W = prs.slide_width
SLIDE_H = prs.slide_height

# ── Helper functions ────────────────────────────────────────────
def set_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_rect(slide, left, top, w, h, fill_color, border_color=None, border_width=Pt(0.75)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    return shape

def thin_line(slide, left, top, width, color, height=Pt(2)):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_link_button(slide, left, top, w, h, text, url, bg_color=SOFT_BLUE, border_color=PRIMARY_BLUE, text_color=PRIMARY_BLUE):
    rect = add_rect(slide, left, top, w, h, bg_color, border_color, Pt(1))
    rect.click_action.hyperlink.address = url
    txBox = slide.shapes.add_textbox(left, top + Pt(2), w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = text_color
    run.font.name = "Segoe UI"
    run.hyperlink.address = url
    return rect

def tb(slide, left, top, w, h, text, size=12, bold=False, color=CHARCOAL,
       align=PP_ALIGN.LEFT, font="Segoe UI", italic=False, spacing_after=Pt(2)):
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font
    p.font.italic = italic
    p.alignment = align
    p.space_after = spacing_after
    return txBox

def multi(slide, left, top, w, h, lines, font="Segoe UI"):
    """lines = [(text, size, bold, color, align, italic?), ...]"""
    txBox = slide.shapes.add_textbox(left, top, w, h)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        text, size, bold, color, align = line[:5]
        italic = line[5] if len(line) > 5 else False
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.font.name = font
        p.font.italic = italic
        p.alignment = align
        p.space_after = Pt(3)
    return txBox


# ================================================================
# SLIDE 1 — TITLE
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

# Subtle accent line at top
thin_line(slide, Inches(0), Inches(0), SLIDE_W, PRIMARY_BLUE, Pt(4))

# Title
tb(slide, Inches(2), Inches(2.2), Inches(9.3), Inches(1.0),
   "Music Streaming Dataset Analysis", 36, True, DEEP_BLUE, PP_ALIGN.CENTER)

# Accent line under title
thin_line(slide, Inches(5.5), Inches(3.4), Inches(2.3), PRIMARY_BLUE, Pt(2))

# Subtitle
tb(slide, Inches(2), Inches(3.7), Inches(9.3), Inches(0.6),
   "Understanding the Drivers of Track Popularity", 18, False, MUTED_GREY, PP_ALIGN.CENTER)

# Metadata
tb(slide, Inches(2), Inches(5.1), Inches(9.3), Inches(0.35),
   "Data Analyst Assessment  |  Tech Round 2", 12, False, MUTED_GREY, PP_ALIGN.CENTER)

# Interactive Chatbot Launch Button
add_link_button(slide, Inches(4.3), Inches(5.65), Inches(4.7), Inches(0.55),
                "💬 Launch Interactive Data Explorer Chatbot", "http://localhost:5050")


# ================================================================
# SLIDE 2 — EXECUTIVE SUMMARY
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(8), Inches(0.5),
   "Executive Summary", 22, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Main headline
tb(slide, Inches(0.8), Inches(1.5), Inches(11), Inches(0.5),
   "Artist context matters more than energy alone.", 16, False, CHARCOAL, italic=True)

# Three findings with numbered labels
findings = [
    ("01", "Energy shows little to no overall association with popularity.",
     "Across 9,637 tracks, there is no consistent relationship between energy level and popularity score."),
    ("02", "The energy\u2013popularity relationship changes substantially across genres.",
     "In some genres (e.g. j-pop, k-pop), higher energy is associated with higher popularity. In others (e.g. emo, garage), the opposite pattern appears."),
    ("03", "Superstar tracks show no meaningful regional concentration.",
     "Superstar-tier artists are distributed proportionally across all five market regions."),
]

for i, (num, headline, detail) in enumerate(findings):
    y = Inches(2.3) + Inches(i * 1.45)
    
    # Light background stripe for each finding
    add_rect(slide, Inches(0.8), y, Inches(11.7), Inches(1.2), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))
    
    # Number
    tb(slide, Inches(1.1), y + Inches(0.15), Inches(0.5), Inches(0.4),
       num, 20, True, PRIMARY_BLUE)
    
    # Headline
    tb(slide, Inches(1.7), y + Inches(0.12), Inches(10.5), Inches(0.35),
       headline, 13, True, CHARCOAL)
    
    # Detail
    tb(slide, Inches(1.7), y + Inches(0.55), Inches(10.5), Inches(0.55),
       detail, 10.5, False, MUTED_GREY)

# Confidence statement
tb(slide, Inches(0.8), Inches(6.8), Inches(11), Inches(0.3),
   "Overall confidence: Moderate to High  |  Based on cleaned dataset of 10,058 tracks",
   10, False, MUTED_GREY, italic=True)


# ================================================================
# SLIDE 3 — DATA QUALITY & CLEANING
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(8), Inches(0.5),
   "Data Quality & Cleaning", 22, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Table header
header_y = Inches(1.5)
col1_x = Inches(0.8)
col2_x = Inches(4.8)
col3_x = Inches(8.8)
col1_w = Inches(3.8)
col2_w = Inches(3.8)
col3_w = Inches(3.8)
row_h = Inches(0.85)

# Header row
add_rect(slide, col1_x, header_y, col1_w, Inches(0.4), SOFT_BLUE, LIGHT_BORDER, Pt(0.5))
add_rect(slide, col2_x, header_y, col2_w, Inches(0.4), SOFT_BLUE, LIGHT_BORDER, Pt(0.5))
add_rect(slide, col3_x, header_y, col3_w, Inches(0.4), SOFT_BLUE, LIGHT_BORDER, Pt(0.5))

tb(slide, col1_x + Inches(0.1), header_y + Inches(0.05), col1_w, Inches(0.3),
   "Issue", 11, True, DEEP_BLUE)
tb(slide, col2_x + Inches(0.1), header_y + Inches(0.05), col2_w, Inches(0.3),
   "What We Did", 11, True, DEEP_BLUE)
tb(slide, col3_x + Inches(0.1), header_y + Inches(0.05), col3_w, Inches(0.3),
   "Why", 11, True, DEEP_BLUE)

# Table rows
rows = [
    ("Inconsistent category casing\n(e.g. 'Europe' vs 'EUROPE')",
     "Standardised all category names\nto consistent title case",
     "Prevented the same group being\ncounted as two separate groups"),
    ("250 non-numeric release years\n(e.g. 'late 90s', 'unknown')",
     "Treated as missing for\nyear-based analysis",
     "Avoided inventing dates that\nwould distort temporal patterns"),
    ("74 exact duplicate rows",
     "Removed exact duplicates;\nkept context-different entries",
     "Same row appearing twice adds\nno information"),
    ("20 invalid popularity values\n(negative or above 100)",
     "Set to missing",
     "Values outside documented\n0\u2013100 range are unreliable"),
    ("4 negative track durations",
     "Set to missing",
     "Physical impossibility indicates\ndata-entry error"),
]

for i, (issue, action, reason) in enumerate(rows):
    y = header_y + Inches(0.4) + i * row_h
    bg = WHITE if i % 2 == 0 else VLIGHT_BLUE
    add_rect(slide, col1_x, y, col1_w, row_h, bg, LIGHT_BORDER, Pt(0.3))
    add_rect(slide, col2_x, y, col2_w, row_h, bg, LIGHT_BORDER, Pt(0.3))
    add_rect(slide, col3_x, y, col3_w, row_h, bg, LIGHT_BORDER, Pt(0.3))
    
    tb(slide, col1_x + Inches(0.1), y + Inches(0.08), col1_w - Inches(0.2), row_h - Inches(0.1),
       issue, 9.5, False, CHARCOAL)
    tb(slide, col2_x + Inches(0.1), y + Inches(0.08), col2_w - Inches(0.2), row_h - Inches(0.1),
       action, 9.5, False, CHARCOAL)
    tb(slide, col3_x + Inches(0.1), y + Inches(0.08), col3_w - Inches(0.2), row_h - Inches(0.1),
       reason, 9.5, False, MUTED_GREY)

# Summary line
tb(slide, Inches(0.8), Inches(6.8), Inches(12), Inches(0.3),
   f"Cleaned dataset: 10,058 rows  |  9,637 with valid popularity scores  |  Every exclusion has a documented rationale",
   10, False, MUTED_GREY, italic=True)


# ================================================================
# SLIDE 4 — HOW THE ANALYSIS WORKED
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(8), Inches(0.5),
   "How the Analysis Worked", 22, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Process flow — 5 steps
steps = [
    ("Data", "10,132 tracks\nacross 114 genres"),
    ("Clean", "Standardise categories\nand remove clear errors"),
    ("Outcome", "Popularity score\n(0\u2013100 scale)"),
    ("Test", "Energy/popularity\n+ regional Superstar"),
    ("Interpret", "Strength of evidence\n+ limitations"),
]

step_w = Inches(2.0)
gap = Inches(0.35)
arrow_w = Inches(0.25)
total = len(steps) * step_w + (len(steps) - 1) * gap
start_x = (SLIDE_W - total) // 2

for i, (title, desc) in enumerate(steps):
    x = start_x + i * (step_w + gap)
    y = Inches(2.0)
    
    # Step box
    add_rect(slide, x, y, step_w, Inches(1.8), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))
    
    # Step label
    tb(slide, x, y + Inches(0.15), step_w, Inches(0.3),
       f"Step {i+1}", 9, False, PRIMARY_BLUE, PP_ALIGN.CENTER)
    
    # Step title
    tb(slide, x, y + Inches(0.4), step_w, Inches(0.35),
       title, 14, True, DEEP_BLUE, PP_ALIGN.CENTER)
    
    # Step description
    tb(slide, x + Inches(0.15), y + Inches(0.85), step_w - Inches(0.3), Inches(0.8),
       desc, 10, False, MUTED_GREY, PP_ALIGN.CENTER)
    
    # Arrow between steps
    if i < len(steps) - 1:
        arrow_x = x + step_w + Inches(0.05)
        tb(slide, arrow_x, y + Inches(0.65), Inches(0.25), Inches(0.3),
           "\u2192", 18, False, PRIMARY_BLUE, PP_ALIGN.CENTER)

# Key principle box
add_rect(slide, Inches(1.5), Inches(4.5), Inches(10.3), Inches(0.55), SOFT_MAGENTA, LIGHT_BORDER, Pt(0.5))
tb(slide, Inches(1.8), Inches(4.58), Inches(9.8), Inches(0.35),
   "Important:  Association does not mean causation. All findings in this analysis describe observed relationships, not proven causes.",
   10.5, False, CHARCOAL, PP_ALIGN.LEFT, italic=True)

# Hypotheses tested
multi(slide, Inches(0.8), Inches(5.5), Inches(11.5), Inches(1.5), [
    ("Hypotheses Tested", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 4, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("H1:  Does energy relate to popularity? Does this vary by genre?", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    ("H2:  Are Superstar artists concentrated in particular regions?", 11, False, CHARCOAL, PP_ALIGN.LEFT),
])


# ================================================================
# SLIDE 5 — H1: ENERGY & POPULARITY
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.5),
   "Energy has little to no overall association with popularity", 20, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Chart
chart_path = os.path.join(CHART_DIR, "h1_energy_popularity.png")
slide.shapes.add_picture(chart_path, Inches(0.5), Inches(1.4), Inches(7.5), Inches(4.2))

# Evidence card on the right
add_rect(slide, Inches(8.3), Inches(1.4), Inches(4.5), Inches(3.0), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))

h1 = R["h1_overall"]
multi(slide, Inches(8.6), Inches(1.5), Inches(4.0), Inches(2.8), [
    ("Evidence", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 6, False, MUTED_GREY, PP_ALIGN.LEFT),
    (f"Correlation:  r \u2248 0  (r = {h1['pearson_r']})", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    (f"p-value:  {h1['pearson_p']}  (not significant)", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    (f"Sample:  {h1['n']:,} tracks", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    ("", 6, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Interpretation", 12, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("Higher energy does not consistently correspond to higher popularity.", 10.5, False, MUTED_GREY, PP_ALIGN.LEFT),
])

# Small technical note
add_rect(slide, Inches(8.3), Inches(4.6), Inches(4.5), Inches(0.6), WHITE, LIGHT_BORDER, Pt(0.3))
tb(slide, Inches(8.5), Inches(4.65), Inches(4.1), Inches(0.5),
   "The relationship is neither statistically significant nor practically meaningful.",
   9, False, MUTED_GREY, italic=True)

# Bottom interpretation
multi(slide, Inches(0.8), Inches(6.0), Inches(11.5), Inches(1.0), [
    ("Mean popularity is remarkably flat across all energy levels (~26\u201335). Low-energy and high-energy tracks achieve similar popularity scores. However, the story changes when we look within individual genres.",
     10.5, False, MUTED_GREY, PP_ALIGN.LEFT),
])

# Footnote
tb(slide, Inches(0.8), Inches(6.9), Inches(11), Inches(0.3),
   "Correlation measures the strength of association between two variables; it does not imply that one causes the other.",
   8.5, False, MUTED_GREY, italic=True)


# ================================================================
# SLIDE 6 — H1: GENRE CONTEXT
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.5),
   "The overall result changes when we look within genres", 20, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Chart
chart_path = os.path.join(CHART_DIR, "h1_genre_correlations.png")
slide.shapes.add_picture(chart_path, Inches(0.3), Inches(1.3), Inches(7.0), Inches(4.3))

# Right panel with findings
add_rect(slide, Inches(7.6), Inches(1.3), Inches(5.2), Inches(4.3), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))

sig_genres = R["h1_significant_genres"]
pos_genres = [g for g in sig_genres if g["pearson_r"] > 0]
neg_genres = [g for g in sig_genres if g["pearson_r"] < 0]

lines = [
    ("What the data shows", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 5, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Some genres show a positive relationship between energy and popularity, while others show the opposite.", 10.5, False, CHARCOAL, PP_ALIGN.LEFT),
    ("", 6, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Higher energy \u2192 higher popularity", 11, True, PRIMARY_BLUE, PP_ALIGN.LEFT),
]

for g in pos_genres[:4]:
    lines.append((f"  {g['genre']}  (r = {g['pearson_r']:+.3f})", 10, False, CHARCOAL, PP_ALIGN.LEFT))

lines.append(("", 5, False, MUTED_GREY, PP_ALIGN.LEFT))
lines.append(("Higher energy \u2192 lower popularity", 11, True, MAGENTA, PP_ALIGN.LEFT))

for g in neg_genres[-4:]:
    lines.append((f"  {g['genre']}  (r = {g['pearson_r']:+.3f})", 10, False, CHARCOAL, PP_ALIGN.LEFT))

multi(slide, Inches(7.9), Inches(1.4), Inches(4.7), Inches(4.1), lines)

# Bottom interpretation
multi(slide, Inches(0.8), Inches(5.9), Inches(11.5), Inches(0.8), [
    ("24 of 114 genres show nominally significant correlations (p < 0.05); results are exploratory due to multiple testing. Genre context appears to be essential \u2014 the relationship between energy and popularity depends on the musical setting.",
     10.5, False, MUTED_GREY, PP_ALIGN.LEFT),
])

# Footnote
tb(slide, Inches(0.8), Inches(6.9), Inches(11), Inches(0.3),
   "Exploratory analysis: testing 114 genres increases the chance of false positives. Individual genre results should be interpreted with caution.",
   8.5, False, MUTED_GREY, italic=True)


# ================================================================
# SLIDE 7 — H2: SUPERSTAR DISTRIBUTION
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.5),
   "Superstar tracks are broadly proportional across regions", 20, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# Chart
chart_path = os.path.join(CHART_DIR, "h2_superstar_regions.png")
slide.shapes.add_picture(chart_path, Inches(0.5), Inches(1.4), Inches(6.5), Inches(3.3))

# Evidence card on the right
add_rect(slide, Inches(7.3), Inches(1.4), Inches(5.5), Inches(3.3), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))

h2 = R["h2_chi2"]
sp = R["h2_superstar_pct"]

multi(slide, Inches(7.6), Inches(1.5), Inches(5.0), Inches(3.1), [
    ("Evidence", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 5, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Range across regions:  3.9% \u2013 5.3%", 11, True, CHARCOAL, PP_ALIGN.LEFT),
    ("", 4, False, MUTED_GREY, PP_ALIGN.LEFT),
    (f"p-value:  {h2['p_value']}  (not significant)", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    (f"Effect size:  Cramer\u2019s V = {h2['cramers_v']}  (negligible)", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    (f"Sample:  {h2['n']:,} tracks", 11, False, CHARCOAL, PP_ALIGN.LEFT),
    ("", 6, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Interpretation", 12, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("The regional differences are small and consistent with random variation.", 10.5, False, MUTED_GREY, PP_ALIGN.LEFT),
])

# Bottom interpretation
multi(slide, Inches(0.8), Inches(5.2), Inches(11.5), Inches(1.0), [
    ("No evidence of regional concentration.  Superstar-tier artists appear across all five market regions at similar rates. The slight variation (3.9%\u20135.3%) is well within what we would expect from chance.", 10.5, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("", 4, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("Note: Middle East & Africa has a smaller sample (588 tracks), which increases uncertainty for that region.", 9.5, False, MUTED_GREY, PP_ALIGN.LEFT, True),
])

# Footnote
tb(slide, Inches(0.8), Inches(6.9), Inches(11), Inches(0.3),
   "Sample: 9,436 tracks with complete popularity, market region, and artist tier data (622 tracks excluded due to missing values).",
   8.5, False, MUTED_GREY, italic=True)


# ================================================================
# SLIDE 8 — CONCLUSIONS, CONFIDENCE & LIMITATIONS
# ================================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, WHITE)

tb(slide, Inches(0.8), Inches(0.5), Inches(8), Inches(0.5),
   "Conclusions, Confidence & Limitations", 22, True, DEEP_BLUE)
thin_line(slide, Inches(0.8), Inches(1.05), Inches(1.8), PRIMARY_BLUE, Pt(2))

# ── LEFT: What the data suggests ──
add_rect(slide, Inches(0.8), Inches(1.4), Inches(5.5), Inches(2.7), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))
multi(slide, Inches(1.1), Inches(1.5), Inches(5.0), Inches(2.5), [
    ("What the Data Suggests", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 4, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("\u2022  Artist tier shows the strongest observed association\n   among the variables examined (Superstar: 76.1 vs Emerging: 11.9).", 10.5, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Energy alone has little overall association\n   with popularity.", 10.5, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Genre context changes the energy\u2013popularity\n   relationship.", 10.5, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Superstar tracks show no meaningful\n   regional concentration.", 10.5, False, CHARCOAL, PP_ALIGN.LEFT),
])

# ── RIGHT: How confident are we? ──
add_rect(slide, Inches(6.5), Inches(1.4), Inches(6.0), Inches(2.7), VLIGHT_BLUE, LIGHT_BORDER, Pt(0.5))
multi(slide, Inches(6.8), Inches(1.5), Inches(5.5), Inches(2.5), [
    ("How Confident Are We?", 13, True, DEEP_BLUE, PP_ALIGN.LEFT),
    ("", 4, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("H1 overall (energy \u2013 popularity):  HIGH", 11, True, CHARCOAL, PP_ALIGN.LEFT),
    ("  Large sample, clear null result.", 9.5, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("", 3, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("H1 genre variation:  MODERATE", 11, True, CHARCOAL, PP_ALIGN.LEFT),
    ("  Meaningful pattern, but ~80\u201390 tracks per genre\n  and multiple testing risk.", 9.5, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("", 3, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("H2 regional distribution:  HIGH", 11, True, CHARCOAL, PP_ALIGN.LEFT),
    ("  Large sample, p = 0.785, negligible effect size.", 9.5, False, MUTED_GREY, PP_ALIGN.LEFT),
])

# ── BOTTOM LEFT: Artist Tier chart ──
chart_path = os.path.join(CHART_DIR, "popularity_by_tier.png")
slide.shapes.add_picture(chart_path, Inches(0.5), Inches(4.4), Inches(5.5), Inches(2.8))

# ── BOTTOM RIGHT: What we cannot claim + final message ──
add_rect(slide, Inches(6.5), Inches(4.4), Inches(6.0), Inches(1.7), WHITE, LIGHT_BORDER, Pt(0.5))
multi(slide, Inches(6.8), Inches(4.5), Inches(5.5), Inches(1.5), [
    ("What We Cannot Claim", 12, True, MAGENTA, PP_ALIGN.LEFT),
    ("", 3, False, MUTED_GREY, PP_ALIGN.LEFT),
    ("\u2022  We cannot establish causation from this data.", 10, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Popularity is observational and snapshot-based.", 10, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Genre, platform and playlist context may\n   independently influence results.", 10, False, CHARCOAL, PP_ALIGN.LEFT),
    ("\u2022  Multiple genre tests require caution when\n   interpreting individual correlations.", 10, False, CHARCOAL, PP_ALIGN.LEFT),
])

# Final analytical message
add_rect(slide, Inches(6.5), Inches(6.2), Inches(6.0), Inches(0.55), SOFT_MAGENTA, LIGHT_BORDER, Pt(0.5))
tb(slide, Inches(6.7), Inches(6.25), Inches(5.6), Inches(0.45),
   "The dataset suggests that artist-level context is more strongly associated with popularity than energy alone.",
   9.5, False, CHARCOAL, italic=True)

# Interactive Chatbot Companion Button
add_link_button(slide, Inches(6.5), Inches(6.82), Inches(6.0), Inches(0.48),
                "💬 Interactive Chatbot Companion — Query 10,058 Rows Live", "http://localhost:5050")



# ================================================================
# SAVE
# ================================================================
output_path = "Music_Streaming_Analysis.pptx"
try:
    prs.save(output_path)
    print(f"[DONE] Revised PowerPoint saved to: {output_path}")
except PermissionError:
    output_path = "Music_Streaming_Analysis_Interactive.pptx"
    prs.save(output_path)
    print(f"[NOTE] Primary file locked by PowerPoint. Saved to: {output_path}")

print(f"  Slides: {len(prs.slides)}")
print(f"  File size: {os.path.getsize(output_path):,} bytes")

