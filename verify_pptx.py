"""
Phase 8 Final: Verify PPT content integrity
"""
from pptx import Presentation
from pptx.util import Inches
import os

pptx_path = "Music_Streaming_Analysis.pptx"
prs = Presentation(pptx_path)

print(f"PPT File: {pptx_path}")
print(f"File size: {os.path.getsize(pptx_path):,} bytes")
print(f"Slide dimensions: {prs.slide_width / 914400:.1f}\" x {prs.slide_height / 914400:.1f}\"")
print(f"Total slides: {len(prs.slides)}")
print()

for i, slide in enumerate(prs.slides):
    shapes = slide.shapes
    text_content = []
    image_count = 0
    
    for shape in shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                if para.text.strip():
                    text_content.append(para.text.strip()[:80])
        if hasattr(shape, 'image'):
            image_count += 1
    
    print(f"--- Slide {i+1} ---")
    print(f"  Shapes: {len(shapes)}")
    print(f"  Images: {image_count}")
    print(f"  Text blocks ({len(text_content)}):")
    for t in text_content[:5]:
        print(f"    \"{t}\"")
    if len(text_content) > 5:
        print(f"    ... and {len(text_content) - 5} more")
    print()

# Verify key content
all_text = []
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                all_text.append(para.text)

full_text = " ".join(all_text)

checks = [
    ("Association", "Association != Causation mentioned"),
    ("r = 0.0109", "H1 Pearson r correct"),
    ("p = 0.785", "H2 chi-square p correct"),
    ("Cramer", "Cramer's V mentioned"),
    ("76.1", "Superstar popularity correct"),
    ("11.9", "Emerging popularity correct"),
    ("j-pop", "Top positive genre mentioned"),
    ("emo", "Top negative genre mentioned"),
    ("250", "Vague release years count"),
    ("74", "Duplicates removed count"),
    ("10,132", "Raw row count"),
    ("9,637", "Analysis sample size"),
    ("Confidence", "Confidence assessment included"),
    ("Limitation", "Limitations section included"),
]

print("=" * 50)
print("CONTENT VERIFICATION")
print("=" * 50)
all_pass = True
for keyword, description in checks:
    found = keyword.lower() in full_text.lower()
    status = "PASS" if found else "FAIL"
    if not found:
        all_pass = False
    print(f"  [{status}] {description} ('{keyword}')")

print(f"\nOverall: {'ALL CHECKS PASSED' if all_pass else 'SOME CHECKS FAILED'}")
