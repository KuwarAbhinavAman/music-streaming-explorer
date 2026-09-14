"""
Verify revised PPT against the brief requirements.
"""
from pptx import Presentation
import os

pptx_path = "Music_Streaming_Analysis.pptx"
prs = Presentation(pptx_path)

print(f"File: {pptx_path}  |  Size: {os.path.getsize(pptx_path):,} bytes  |  Slides: {len(prs.slides)}")
print()

# Extract all text
all_text = []
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                all_text.append(para.text)
full = " ".join(all_text).lower()

print("=" * 60)
print("CONTENT CHECKS")
print("=" * 60)

checks = [
    # Correct language
    ("association", "Uses 'association' not 'effect'"),
    ("does not mean causation", "Association ≠ causation stated"),
    ("observed", "'Observed' association language"),
    ("little to no", "'Little to no' wording used"),
    ("cannot establish causation", "Causation disclaimer present"),
    
    # Key numbers
    ("0.0109", "H1 Pearson r correct"),
    ("0.285", "H1 p-value correct"),
    ("9,637", "H1 sample size correct"),
    ("0.785", "H2 p-value correct"),
    ("0.0168", "H2 Cramer's V correct"),
    ("9,436", "H2 sample size correct"),
    ("10,058", "Cleaned row count"),
    ("10,132", "Raw row count"),
    ("76.1", "Superstar mean popularity"),
    ("11.9", "Emerging mean popularity"),
    
    # Genre findings
    ("j-pop", "Top positive genre"),
    ("emo", "Top negative genre"),
    ("24 of 114", "Qualified genre findings"),
    ("false positive", "Multiple testing caveat"),
    
    # Correct framing
    ("snapshot", "Snapshot nature acknowledged"),
    ("middle east", "ME&A small sample noted"),
    
    # Things that should NOT be present
]

all_pass = True
for keyword, description in checks:
    found = keyword.lower() in full
    status = "PASS" if found else "FAIL"
    if not found:
        all_pass = False
    print(f"  [{status}] {description}")

# Negative checks — things that SHOULD NOT be present
print()
print("=" * 60)
print("NEGATIVE CHECKS (should NOT appear)")
print("=" * 60)

neg_checks = [
    ("strongest predictor", "No 'predictor' language"),
    ("driver", "No 'driver' language"),
    ("null hypothesis", "No 'null hypothesis'"),
    ("variance explained", "No 'variance explained'"),
    ("music industry research", "No unsourced external claims"),
    ("marketing, branding", "No unsourced external claims (v2)"),
    ("logo", "No logo references"),
    ("heteroscedasticity", "No advanced stats jargon"),
    ("multicollinearity", "No advanced stats jargon (v2)"),
]

for keyword, description in neg_checks:
    found = keyword.lower() in full
    status = "PASS" if not found else "FAIL"
    if found:
        all_pass = False
    print(f"  [{status}] {description}")

print(f"\nOverall: {'ALL CHECKS PASSED' if all_pass else 'SOME CHECKS FAILED — review above'}")

# Show slide structure
print()
print("=" * 60)
print("SLIDE STRUCTURE")
print("=" * 60)
for i, slide in enumerate(prs.slides):
    shapes = slide.shapes
    images = sum(1 for s in shapes if hasattr(s, 'image'))
    texts = []
    for s in shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip() and len(p.text.strip()) > 5:
                    texts.append(p.text.strip()[:90])
    
    print(f"\n  Slide {i+1}: {len(shapes)} shapes, {images} images")
    for t in texts[:3]:
        print(f"    \"{t}\"")
